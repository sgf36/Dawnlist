"""Tests for app.export.letters — letter .docx generation."""
import pytest

from docx import Document
from docx.shared import Mm

from app.core import db
from app.export.letters import (
    LetterSpec, build_letters_docx, A4_WIDTH_TWIPS, A4_HEIGHT_TWIPS,
)


@pytest.fixture()
def conn(tmp_path):
    c = db.connect(tmp_path / "t.sqlite3")
    db.migrate(c)
    yield c
    c.close()


def test_single_letter(tmp_path):
    out = tmp_path / "letters.docx"
    specs = [LetterSpec(
        recipient_name="Paul Blackmore",
        recipient_title="Mr",
        company="Hilton Hotels",
        address_lines=["Maple Court", "Central Park Avenue",
                        "Repton Park", "IG8 8GZ"],
        body="Dear Mr Blackmore,\n\nI am writing to you regarding...",
        opportunity_id=1,
    )]
    build_letters_docx(specs, out)
    assert out.exists()

    doc = Document(str(out))
    sec = doc.sections[0]
    assert round(sec.page_width / 635) == A4_WIDTH_TWIPS
    assert round(sec.page_height / 635) == A4_HEIGHT_TWIPS
    assert abs(sec.left_margin - Mm(25)) < 1000  # EMU rounding tolerance


def test_multiple_letters_one_page_each(tmp_path):
    out = tmp_path / "letters.docx"
    specs = [
        LetterSpec("Alice Smith", "", "Savills",
                   ["1 Mayfair Place", "London", "W1K 1AB"],
                   "Dear Alice,\n\nThank you.", 1),
        LetterSpec("Bob Jones", "Mr", "CBRE",
                   ["10 South Place", "London", "EC2M 7EB"],
                   "Dear Mr Jones,\n\nI hope.", 2),
    ]
    build_letters_docx(specs, out)
    doc = Document(str(out))
    text = "\n".join(p.text for p in doc.paragraphs)
    assert "Alice Smith" in text
    assert "Bob Jones" in text


def test_empty_list_raises():
    import tempfile
    with tempfile.TemporaryDirectory() as d:
        from pathlib import Path
        with pytest.raises(ValueError, match="No letters"):
            build_letters_docx([], Path(d) / "empty.docx")


def test_migration_adds_contact_fields(conn):
    cols = {row[1] for row in
            conn.execute("PRAGMA table_info(contacts)").fetchall()}
    assert "first_name" in cols
    assert "last_name" in cols
    assert "job_title" in cols
    assert "phone" in cols
    assert "mailing_address" in cols


def test_migration_adds_clickup_task_id(conn):
    cols = {row[1] for row in
            conn.execute("PRAGMA table_info(opportunities)").fetchall()}
    assert "clickup_task_id" in cols
