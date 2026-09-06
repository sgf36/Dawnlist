"""Reading a CV corpus off disk, including the scanned-PDF case."""
from pathlib import Path

import pytest

from app.onboarding.extract import (SCAN_CHARS_PER_PAGE, extract_corpus,
                                    extract_one, gather)

CV_TEXT = ("Spencer Fields\n\nAsset Manager, Acme Hotels, 2019 to 2023\n"
           "Identified GBP 4.2m in operational savings opportunities.\n"
           "Supported the disposal of three assets.\n") * 6


def make_docx(path: Path, *, with_table: bool = False) -> Path:
    import docx
    d = docx.Document()
    for line in CV_TEXT.splitlines():
        d.add_paragraph(line)
    if with_table:
        t = d.add_table(rows=2, cols=2)
        t.cell(0, 0).text = "Employer"
        t.cell(0, 1).text = "Dates"
        t.cell(1, 0).text = "Rocco Forte Hotels"
        t.cell(1, 1).text = "2016 to 2019"
    d.save(str(path))
    return path


def make_text_pdf(path: Path) -> Path:
    """A real PDF with real text, drawn with Qt so the test needs no extra dep."""
    from PySide6.QtGui import QPageSize, QPdfWriter, QPainter
    from PySide6.QtWidgets import QApplication

    QApplication.instance() or QApplication([])
    writer = QPdfWriter(str(path))
    writer.setPageSize(QPageSize(QPageSize.A4))
    painter = QPainter(writer)
    y = 400
    for line in CV_TEXT.splitlines():
        painter.drawText(400, y, line)
        y += 300
    painter.end()
    return path


def make_scanned_pdf(path: Path, pages: int = 2) -> Path:
    """Pages, but no extractable text — what a scan looks like to a reader."""
    from pypdf import PdfWriter
    w = PdfWriter()
    for _ in range(pages):
        w.add_blank_page(width=595, height=842)
    with open(path, "wb") as fh:
        w.write(fh)
    return path


# -- the happy paths --------------------------------------------------------
def test_a_docx_is_read(tmp_path):
    doc, failure = extract_one(make_docx(tmp_path / "cv.docx"))
    assert failure is None
    assert "Asset Manager" in doc.text and doc.name == "cv.docx"


def test_docx_tables_are_read_too(tmp_path):
    """CVs routinely put employers and dates in a table, and a paragraph-only
    read silently drops every one of them."""
    doc, _ = extract_one(make_docx(tmp_path / "cv.docx", with_table=True))
    assert "Rocco Forte Hotels" in doc.text
    assert "2016 to 2019" in doc.text


def test_a_text_pdf_is_read(tmp_path):
    doc, failure = extract_one(make_text_pdf(tmp_path / "cv.pdf"))
    assert failure is None, f"unexpected failure: {failure}"
    assert "Spencer Fields" in doc.text


def test_plain_text_is_read(tmp_path):
    p = tmp_path / "cv.txt"
    p.write_text(CV_TEXT, encoding="utf-8")
    doc, failure = extract_one(p)
    assert failure is None and "Identified" in doc.text


def test_a_cp1252_file_is_decoded(tmp_path):
    p = tmp_path / "cv.txt"
    p.write_bytes(("Rocco Forte — Hôtel de Rome\n" + CV_TEXT).encode("cp1252"))
    doc, failure = extract_one(p)
    assert failure is None and "Hôtel de Rome" in doc.text


def test_paragraph_structure_survives_cleaning(tmp_path):
    p = tmp_path / "cv.txt"
    p.write_text("Role one\n\n\n\nRole two\n" + CV_TEXT, encoding="utf-8")
    doc, _ = extract_one(p)
    assert "Role one\n\nRole two" in doc.text, (
        "blank lines between roles are structure, not noise")


# -- the scanned PDF, which is the whole point ------------------------------
def test_a_scanned_pdf_is_diagnosed_by_name_not_treated_as_empty(tmp_path):
    doc, failure = extract_one(make_scanned_pdf(tmp_path / "scan.pdf"))
    assert doc is None
    assert failure.name == "scan.pdf"
    assert "scanned PDF" in failure.reason
    assert "add the .docx instead" in failure.reason, (
        "the message must tell the user what to do next")


def test_a_scanned_pdf_never_joins_the_corpus_silently(tmp_path):
    result = extract_corpus([make_docx(tmp_path / "cv.docx"),
                             make_scanned_pdf(tmp_path / "scan.pdf")])
    assert len(result.corpus.documents) == 1
    assert any("scan.pdf" in w for w in result.warnings)


# -- failures are named, never swallowed ------------------------------------
def test_an_old_doc_file_says_what_to_do(tmp_path):
    p = tmp_path / "old.doc"
    p.write_bytes(b"\xd0\xcf\x11\xe0legacy")
    doc, failure = extract_one(p)
    assert doc is None and "save as .docx" in failure.reason


def test_an_unsupported_type_is_named(tmp_path):
    p = tmp_path / "photo.png"
    p.write_bytes(b"\x89PNG")
    doc, failure = extract_one(p)
    assert doc is None and "unsupported file type" in failure.reason


def test_a_missing_file_is_reported(tmp_path):
    doc, failure = extract_one(tmp_path / "gone.docx")
    assert doc is None and "no longer exists" in failure.reason


def test_a_corrupt_pdf_is_reported_not_raised(tmp_path):
    p = tmp_path / "broken.pdf"
    p.write_bytes(b"%PDF-1.4 this is not really a pdf")
    doc, failure = extract_one(p)
    assert doc is None and failure.name == "broken.pdf"


def test_an_empty_file_is_reported(tmp_path):
    p = tmp_path / "empty.txt"
    p.write_text("", encoding="utf-8")
    doc, failure = extract_one(p)
    assert doc is None and "no text" in failure.reason


def test_one_bad_file_never_stops_the_others(tmp_path):
    (tmp_path / "junk.png").write_bytes(b"\x89PNG")
    result = extract_corpus([tmp_path / "junk.png",
                             make_docx(tmp_path / "a.docx"),
                             make_docx(tmp_path / "b.docx")])
    assert len(result.corpus.documents) == 2
    assert len(result.failures) == 1


# -- gathering --------------------------------------------------------------
def test_gather_finds_candidates_and_skips_word_lock_files(tmp_path):
    make_docx(tmp_path / "cv.docx")
    (tmp_path / "~$cv.docx").write_bytes(b"lock")
    (tmp_path / "notes.txt").write_text("x", encoding="utf-8")
    (tmp_path / "photo.png").write_bytes(b"\x89PNG")
    names = [p.name for p in gather(tmp_path)]
    assert names == ["cv.docx", "notes.txt"]


def test_gather_recurses(tmp_path):
    (tmp_path / "old").mkdir()
    make_docx(tmp_path / "old" / "cv2019.docx")
    make_docx(tmp_path / "cv.docx")
    assert len(gather(tmp_path)) == 2


# -- the corpus warnings still apply ----------------------------------------
def test_corpus_warnings_are_merged_with_extraction_failures(tmp_path):
    result = extract_corpus([make_docx(tmp_path / "cv.docx"),
                             tmp_path / "gone.pdf"])
    warnings = result.warnings
    assert any("gone.pdf" in w for w in warnings)
    assert any("Only one CV version" in w for w in warnings)
