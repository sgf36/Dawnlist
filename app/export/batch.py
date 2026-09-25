"""Generate an EasyPost batch-import workbook from letter specs.

Reads the bundled batch_template.xlsx template and fills one row per letter.
The to_country and predefined_package values are resolved programmatically
from the template's own reference sheets — never typed as literals (the
em-dash in those strings has caused silent rejections before).
"""
from __future__ import annotations

from pathlib import Path

from openpyxl import load_workbook

from app.export.letters import LetterSpec

RESOURCES = Path(__file__).resolve().parent.parent / "resources"
TEMPLATE = RESOURCES / "batch_template.xlsx"

LETTER_WEIGHT_OZ = 3.5


def _resolve_from_sheet(wb, sheet_name: str, search: str) -> str:
    """Find the exact value from a reference sheet by substring match."""
    ws = wb[sheet_name]
    for row in ws.iter_rows(min_row=2, values_only=True):
        val = row[0]
        if val and search.lower() in str(val).lower():
            return str(val)
    raise ValueError(
        f"Could not resolve '{search}' from {sheet_name} sheet")


def build_batch_xlsx(letters: list[LetterSpec], out: Path,
                     template: Path | None = None) -> Path:
    """Fill one row per letter into a copy of the batch template."""
    if not letters:
        raise ValueError("No letters for batch")

    template = template or TEMPLATE
    if not template.exists():
        raise FileNotFoundError(f"Batch template not found: {template}")

    wb = load_workbook(str(template))
    country_gb = _resolve_from_sheet(wb, "Countries", "GB")
    package_letter = _resolve_from_sheet(wb, "Packages",
                                         "Royal Mail V3")
    # Narrow to the Letter variant, not LargeLetter
    for row in wb["Packages"].iter_rows(min_row=2, values_only=True):
        val = row[0]
        if val and "Royal Mail V3" in str(val) and "Letter" in str(val):
            if "Large" not in str(val):
                package_letter = str(val)
                break

    ws = wb["Recipients"]
    # Clear the example row if present
    if ws.max_row > 1:
        ws.delete_rows(2, ws.max_row - 1)

    for letter in letters:
        name_parts = []
        if letter.recipient_title:
            name_parts.append(letter.recipient_title)
        if letter.recipient_name:
            name_parts.append(letter.recipient_name)
        to_name = " ".join(name_parts)

        street1 = ""
        street2 = ""
        city = ""
        state = ""
        postcode = ""
        if letter.address_lines:
            if len(letter.address_lines) >= 1:
                street1 = letter.address_lines[0]
            if len(letter.address_lines) >= 3:
                street2 = letter.address_lines[1]
                city = letter.address_lines[-2]
            elif len(letter.address_lines) >= 2:
                city = letter.address_lines[-2] if len(
                    letter.address_lines) > 2 else ""
            postcode = letter.address_lines[-1] if letter.address_lines else ""
            # Heuristic: last line that looks like a UK postcode
            for line in reversed(letter.address_lines):
                stripped = line.strip().upper()
                if len(stripped) >= 5 and " " in stripped and any(
                        c.isdigit() for c in stripped):
                    postcode = stripped
                    break

        ws.append([
            to_name,                # to_name
            letter.company,         # to_company
            street1,                # to_street1
            street2,                # to_street2
            city,                   # to_city
            state,                  # to_state
            postcode,               # to_zip
            country_gb,             # to_country
            "",                     # to_phone
            "",                     # to_email
            "",                     # length
            "",                     # width
            "",                     # height
            LETTER_WEIGHT_OZ,       # weight
            package_letter,         # predefined_package
            str(letter.opportunity_id),  # reference
        ])

    wb.save(str(out))
    return out
