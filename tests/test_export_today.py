"""Tests for app.export.today — do-today .docx and .xlsx generation."""
import pytest

from docx import Document
from docx.shared import Twips

from app.core import db
from app.export.today import (TodaySpec, TodayRow, build_today_docx,
                               build_today_xlsx, MARGIN)


@pytest.fixture()
def conn(tmp_path):
    c = db.connect(tmp_path / "t.sqlite3")
    db.migrate(c)
    yield c
    c.close()


def test_basic_render(tmp_path):
    out = tmp_path / "today.docx"
    spec = TodaySpec(
        title="Thursday 25 September 2026",
        rows=[
            TodayRow("send", "2026-09-25", "Follow-up email", "Alice, Savills",
                     "Email", opportunity_id=1),
            TodayRow("post", "2026-09-25", "Handwritten letter",
                     "Bob, CBRE", "Letter", opportunity_id=2),
        ],
    )
    build_today_docx(spec, out)
    assert out.exists()

    doc = Document(str(out))
    sec = doc.sections[0]
    assert round(sec.page_width / 635) == 16838
    assert round(sec.page_height / 635) == 11906
    assert sec.left_margin == MARGIN


def test_blocked_items_in_footer(tmp_path):
    out = tmp_path / "today.docx"
    spec = TodaySpec(
        title="Test",
        rows=[TodayRow("send", "2026-09-25", "Email", "Alice, Co", "Email")],
        footer=["BLOCKED: CBRE call — no phone number"],
    )
    build_today_docx(spec, out)
    doc = Document(str(out))
    text = "\n".join(p.text for p in doc.paragraphs)
    assert "BLOCKED" in text


def test_empty_bands_omitted(tmp_path):
    out = tmp_path / "today.docx"
    spec = TodaySpec(
        title="Test",
        rows=[TodayRow("send", "2026-09-25", "Email", "A, B", "Email")],
    )
    build_today_docx(spec, out)
    doc = Document(str(out))
    text = "\n".join(p.text for p in doc.paragraphs)
    assert "Letters to post" not in text
    assert "Send now" in text


def test_notes_column_always_empty(tmp_path):
    out = tmp_path / "today.docx"
    spec = TodaySpec(
        title="Test",
        rows=[TodayRow("send", "2026-09-25", "Email", "A, B", "Email")],
    )
    build_today_docx(spec, out)
    doc = Document(str(out))
    for table in doc.tables:
        for row in table.rows[1:]:
            assert row.cells[-1].text == ""


def test_xlsx_basic_render(tmp_path):
    out = tmp_path / "today.xlsx"
    spec = TodaySpec(
        title="Thursday 25 September 2026",
        rows=[
            TodayRow("send", "2026-09-25", "Follow-up email", "Alice, Savills",
                     "Email", opportunity_id=1),
            TodayRow("post", "2026-09-25", "Handwritten letter",
                     "Bob, CBRE", "Letter", opportunity_id=2),
        ],
    )
    build_today_xlsx(spec, out)
    assert out.exists()

    from openpyxl import load_workbook
    wb = load_workbook(str(out))
    ws = wb.active
    assert ws.cell(row=1, column=1).value == "#"
    assert ws.cell(row=1, column=2).value == "Band"
    assert ws.cell(row=2, column=1).value == 1
    assert ws.cell(row=3, column=1).value == 2
    assert ws.auto_filter.ref is not None


def test_xlsx_blocked_in_footer(tmp_path):
    out = tmp_path / "today.xlsx"
    spec = TodaySpec(
        title="Test",
        rows=[TodayRow("send", "2026-09-25", "Email", "Alice, Co", "Email")],
        footer=["BLOCKED: CBRE call — no phone number"],
    )
    build_today_xlsx(spec, out)

    from openpyxl import load_workbook
    wb = load_workbook(str(out))
    ws = wb.active
    found = False
    for row in ws.iter_rows(values_only=True):
        for cell in row:
            if cell and "BLOCKED" in str(cell):
                found = True
    assert found


def test_xlsx_notes_column_empty(tmp_path):
    out = tmp_path / "today.xlsx"
    spec = TodaySpec(
        title="Test",
        rows=[TodayRow("send", "2026-09-25", "Email", "A, B", "Email")],
    )
    build_today_xlsx(spec, out)

    from openpyxl import load_workbook
    wb = load_workbook(str(out))
    ws = wb.active
    for row in ws.iter_rows(min_row=2, values_only=True):
        if row[0] is not None:
            assert row[-1] is None or row[-1] == ""
