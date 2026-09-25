"""Render a do-today task sheet from Dawnlist's database.

A4 landscape, banded tables grouped by channel. Layout matches the reference
implementation (scripts/daily_tasks_docx.py) and the routine spec (F2).

Margins pinned at 900 twips. Verified A4 landscape: 16838 × 11906 twips.

Two output formats: .docx for printing and .xlsx for sorting/filtering.
"""
from __future__ import annotations

import sqlite3
from dataclasses import dataclass, field
from datetime import date
from pathlib import Path

from docx import Document
from docx.shared import Mm, Pt, Twips, RGBColor
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

MARGIN = Twips(900)
COLS = ["#", "Due", "What to do", "Who", "Route", "Done", "Notes"]
COL_WIDTHS_MM = [8, 25, 95, 38, 55, 11, 19]

BANDS = [
    ("send", "Send now", "emails ready in Drafts"),
    ("calls", "Calls", None),
    ("post", "Letters to post", None),
    ("apply", "Apply online", None),
    ("route", "Find a route", None),
]

CHANNEL_TO_BAND = {
    "email": "send",
    "letter": "post",
    "call": "calls",
}


@dataclass
class TodayRow:
    band: str
    due: str
    what: str
    who: str
    route: str
    bold_due: bool = False
    opportunity_id: int | None = None


@dataclass
class TodaySpec:
    title: str
    subtitle: str = ""
    deadlines: list[str] = field(default_factory=list)
    rows: list[TodayRow] = field(default_factory=list)
    footer: list[str] = field(default_factory=list)


def _shade(cell, fill: str) -> None:
    sh = OxmlElement("w:shd")
    sh.set(qn("w:val"), "clear")
    sh.set(qn("w:color"), "auto")
    sh.set(qn("w:fill"), fill)
    cell._tc.get_or_add_tcPr().append(sh)


def _flag(row, tag: str) -> None:
    pr = row._tr.get_or_add_trPr()
    e = OxmlElement(tag)
    e.set(qn("w:val"), "true")
    pr.append(e)


def _borders(tbl) -> None:
    b = OxmlElement("w:tblBorders")
    for k in ("top", "left", "bottom", "right", "insideH", "insideV"):
        e = OxmlElement("w:" + k)
        for attr, val in (("val", "single"), ("sz", "4"),
                          ("space", "0"), ("color", "999999")):
            e.set(qn("w:" + attr), val)
        b.append(e)
    tbl._tbl.tblPr.append(b)


def _fix_grid(tbl, widths: list[int]) -> None:
    for gc, w in zip(
            tbl._tbl.tblGrid.findall(qn("w:gridCol")), widths):
        gc.set(qn("w:w"), str(int(w * 56.6929)))
    lay = OxmlElement("w:tblLayout")
    lay.set(qn("w:type"), "fixed")
    tbl._tbl.tblPr.append(lay)
    for row in tbl.rows:
        for i, w in enumerate(widths):
            row.cells[i].width = Mm(w)


def _text(cell, s: str, bold: bool = False, size: int = 8) -> None:
    cell.text = ""
    p = cell.paragraphs[0]
    lines = (s or "").split("\n")
    for i, ln in enumerate(lines):
        r = p.add_run(ln)
        r.bold = bold
        r.font.size = Pt(size)
        if i < len(lines) - 1:
            r.add_break()


def build_today_docx(spec: TodaySpec, out: Path) -> Path:
    """Render the do-today sheet."""
    doc = Document()
    s = doc.sections[0]
    s.orientation = WD_ORIENT.LANDSCAPE
    s.page_width = Mm(297)
    s.page_height = Mm(210)
    s.left_margin = s.right_margin = MARGIN
    s.top_margin = s.bottom_margin = Mm(12)

    st = doc.styles["Normal"]
    st.font.name = "Arial"
    st.font.size = Pt(9)
    st.element.rPr.rFonts.set(qn("w:eastAsia"), "Arial")
    st.paragraph_format.space_after = Pt(0)

    def para(text="", bold=False, size=9, color=None, after=0, before=0):
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(after)
        p.paragraph_format.space_before = Pt(before)
        p.paragraph_format.keep_with_next = bold
        if text:
            r = p.add_run(text)
            r.bold = bold
            r.font.size = Pt(size)
            if color:
                r.font.color.rgb = RGBColor.from_string(color)
        return p

    def bullets(items, size):
        for b in items:
            p = doc.add_paragraph(style="List Bullet")
            p.paragraph_format.space_after = Pt(0)
            p.add_run(b).font.size = Pt(size)

    para(f"Do today · {spec.title}", bold=True, size=15, after=2)
    if spec.subtitle:
        para(spec.subtitle, size=8, color="555555", after=4)
    if spec.deadlines:
        para("Deadlines that bite", bold=True, size=10, after=1)
        bullets(spec.deadlines, 9)

    n = 0
    for key, title, sub in BANDS:
        band_rows = [r for r in spec.rows if r.band == key]
        if not band_rows:
            continue
        head = f"{title} ({len(band_rows)}"
        if sub:
            head += f"; {sub}"
        head += ")"
        para(head, bold=True, size=11, before=8, after=1)

        t = doc.add_table(rows=1, cols=len(COLS))
        t.alignment = WD_TABLE_ALIGNMENT.LEFT
        t.autofit = False
        _borders(t)
        for i, c in enumerate(COLS):
            _text(t.rows[0].cells[i], c, bold=True)
            _shade(t.rows[0].cells[i], "D9E2F3")
        _flag(t.rows[0], "w:tblHeader")

        for r in band_rows:
            n += 1
            row = t.add_row()
            _flag(row, "w:cantSplit")
            vals = [str(n), r.due, r.what, r.who, r.route]
            for i, v in enumerate(vals):
                _text(row.cells[i], v,
                      bold=(i == 1 and r.bold_due))
            _text(row.cells[5], "☐", size=11)
            _text(row.cells[6], "")
        _fix_grid(t, COL_WIDTHS_MM)

    if spec.footer:
        para("Not actionable today", bold=True, size=10, before=8,
             after=1)
        bullets(spec.footer, 8.5)

    doc.save(str(out))
    _verify_geometry(out)
    return out


def _verify_geometry(path: Path) -> None:
    doc = Document(str(path))
    sec = doc.sections[0]
    w = round(sec.page_width / 635)
    h = round(sec.page_height / 635)
    if (w, h) != (16838, 11906):
        raise AssertionError(
            f"Today .docx geometry {w}×{h}, expected 16838×11906")
    if sec.left_margin != MARGIN or sec.right_margin != MARGIN:
        raise AssertionError("Today .docx margins not 900 twips")


def build_today_xlsx(spec: TodaySpec, out: Path) -> Path:
    """Render the do-today sheet as a filterable .xlsx workbook."""
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

    wb = Workbook()
    ws = wb.active
    ws.title = "Do today"

    header_font = Font(name="Arial", size=9, bold=True)
    header_fill = PatternFill("solid", fgColor="D9E2F3")
    cell_font = Font(name="Arial", size=9)
    thin = Side(style="thin", color="999999")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)

    headers = ["#", "Band", "Due", "What to do", "Who", "Route", "Done",
               "Notes"]
    widths = [5, 14, 12, 50, 30, 14, 8, 20]
    for i, (h, w) in enumerate(zip(headers, widths), 1):
        c = ws.cell(row=1, column=i, value=h)
        c.font = header_font
        c.fill = header_fill
        c.border = border
        c.alignment = Alignment(wrap_text=True)
        ws.column_dimensions[chr(64 + i)].width = w

    n = 0
    for key, title, _sub in BANDS:
        band_rows = [r for r in spec.rows if r.band == key]
        for r in band_rows:
            n += 1
            row_num = n + 1
            vals = [n, title, r.due, r.what, r.who, r.route, "", ""]
            for col, v in enumerate(vals, 1):
                c = ws.cell(row=row_num, column=col, value=v)
                c.font = cell_font
                c.border = border
                c.alignment = Alignment(wrap_text=(col == 4))

    if spec.footer:
        gap_row = n + 2
        ws.cell(row=gap_row, column=1)
        for i, line in enumerate(spec.footer):
            c = ws.cell(row=gap_row + 1 + i, column=3, value=line)
            c.font = Font(name="Arial", size=8, italic=True, color="CC0000")

    ws.auto_filter.ref = f"A1:H{max(n + 1, 2)}"
    ws.freeze_panes = "A2"

    wb.save(str(out))
    return out


def tasks_due_today(conn: sqlite3.Connection,
                    as_of: date | None = None) -> TodaySpec:
    """Build a TodaySpec from tasks due on or before as_of."""
    as_of = as_of or date.today()

    rows = conn.execute("""
        SELECT t.id, t.title, t.due_on, t.status,
               o.id AS opp_id, o.company, o.stage,
               c.name, c.first_name, c.last_name,
               c.mailing_address, c.phone
        FROM tasks t
        JOIN opportunities o ON o.id = t.opportunity_id
        LEFT JOIN contacts c ON c.opportunity_id = o.id
                            AND c.do_not_contact = 0
        WHERE t.status IN ('open', 'waiting')
          AND t.due_on <= ?
        ORDER BY t.due_on, o.company
    """, (as_of.isoformat(),)).fetchall()

    today_rows = []
    footer = []
    for r in rows:
        who = r["name"] or ""
        if r["first_name"] and r["last_name"]:
            who = f"{r['first_name']} {r['last_name']}"
        who_line = f"{who}, {r['company']}" if who else r["company"]

        title_lower = r["title"].lower()
        if "letter" in title_lower:
            band = "post"
            if not r["mailing_address"]:
                footer.append(
                    f"BLOCKED: {r['company']} letter — no postal address")
                continue
        elif "call" in title_lower or "phone" in title_lower:
            band = "calls"
            if not r["phone"]:
                footer.append(
                    f"BLOCKED: {r['company']} call — no phone number")
                continue
        elif "email" in title_lower or "send" in title_lower:
            band = "send"
        elif "appl" in title_lower or "portal" in title_lower:
            band = "apply"
        else:
            band = "route"

        route = ""
        if "email" in title_lower and "letter" in title_lower:
            route = "Email + Letter"
        elif "letter" in title_lower:
            route = "Letter"
        elif "email" in title_lower:
            route = "Email"
        elif "call" in title_lower:
            route = "Phone"
        elif "appl" in title_lower:
            route = "Online"

        today_rows.append(TodayRow(
            band=band,
            due=r["due_on"] or "",
            what=r["title"],
            who=who_line,
            route=route,
            opportunity_id=r["opp_id"],
        ))

    import locale
    try:
        title = as_of.strftime("%A %d %B %Y")
    except ValueError:
        title = as_of.isoformat()

    return TodaySpec(
        title=title,
        rows=today_rows,
        footer=footer,
    )
