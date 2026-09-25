"""Tests for app.export.batch — EasyPost batch .xlsx generation."""
import pytest

from openpyxl import load_workbook

from app.export.batch import build_batch_xlsx, TEMPLATE
from app.export.letters import LetterSpec


def test_batch_fills_from_template(tmp_path):
    out = tmp_path / "batch.xlsx"
    specs = [LetterSpec(
        recipient_name="Paul Blackmore",
        recipient_title="Mr",
        company="Hilton",
        address_lines=["10 Fleet Place", "London", "EC4M 7RB"],
        body="test",
        opportunity_id=42,
    )]
    build_batch_xlsx(specs, out)
    assert out.exists()

    wb = load_workbook(str(out))
    ws = wb["Recipients"]
    assert ws.max_row == 2  # header + 1 data row
    assert ws.cell(2, 1).value == "Mr Paul Blackmore"
    assert ws.cell(2, 2).value == "Hilton"

    country = ws.cell(2, 8).value
    assert "—" in country  # em-dash from reference sheet
    assert "United Kingdom" in country

    package = ws.cell(2, 15).value
    assert "Royal Mail" in package
    assert "Letter" in package

    assert ws.cell(2, 16).value == "42"


def test_batch_rejects_empty():
    import tempfile
    from pathlib import Path
    with tempfile.TemporaryDirectory() as d:
        with pytest.raises(ValueError, match="No letters"):
            build_batch_xlsx([], Path(d) / "empty.xlsx")


@pytest.mark.skipif(not TEMPLATE.exists(),
                    reason="batch_template.xlsx not bundled yet")
def test_template_has_reference_sheets():
    wb = load_workbook(str(TEMPLATE))
    assert "Packages" in wb.sheetnames
    assert "Countries" in wb.sheetnames
